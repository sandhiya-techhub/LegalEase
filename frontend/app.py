import streamlit as st
import requests
from io import BytesIO
from docx import Document
from reportlab.pdfgen import canvas


st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide"
)

st.title("⚖️ LegalEase")
st.subheader("AI-Powered Legal Document Generator")


document_type = st.selectbox(
    "Select Document Type",
    [
        "Employment Contract",
        "Non-Disclosure Agreement (NDA)",
        "Residential Lease Agreement"
    ]
)


parties = st.text_area(
    "Parties Involved",
    placeholder="Example: Employer and Employee"
)


terms = st.text_area(
    "Terms & Conditions",
    placeholder="Enter the important terms and conditions"
)


dates = st.text_input(
    "Effective Date",
    placeholder="Example: 24-09-2026"
)


if st.button("Generate Legal Document"):

    if not parties or not terms or not dates:
        st.warning("Please fill in all the required fields.")

    else:
        try:
            response = requests.post(
                "http://127.0.0.1:8000/generate",
                json={
                    "document_type": document_type,
                    "parties": parties,
                    "terms": terms,
                    "dates": dates
                }
            )

            if response.status_code == 200:
                generated_text = response.json()["document"]

                st.session_state.generated_text = generated_text

                st.success("Legal document generated successfully!")

            else:
                st.error("Unable to generate the document.")

        except Exception as e:
            st.error(f"Connection error: {e}")


if "generated_text" in st.session_state:

    st.subheader("Generated Legal Document")

    edited_text = st.text_area(
        "Edit your document",
        value=st.session_state.generated_text,
        height=400
    )

    st.session_state.generated_text = edited_text

    st.subheader("Download Document")

    # TXT Download
    st.download_button(
        label="📄 Download TXT",
        data=edited_text,
        file_name="LegalEase_Document.txt",
        mime="text/plain"
    )

    # DOCX Download
    doc = Document()
    doc.add_heading("LegalEase - Legal Document", level=1)

    for paragraph in edited_text.split("\n"):
        doc.add_paragraph(paragraph)

    docx_file = BytesIO()
    doc.save(docx_file)
    docx_file.seek(0)

    st.download_button(
        label="📝 Download DOCX",
        data=docx_file,
        file_name="LegalEase_Document.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    # PDF Download
    pdf_file = BytesIO()

    pdf = canvas.Canvas(pdf_file)
    pdf.setTitle("LegalEase Legal Document")

    text_object = pdf.beginText(50, 800)
    text_object.setFont("Helvetica", 11)

    for line in edited_text.split("\n"):
        if text_object.getY() < 50:
            pdf.drawText(text_object)
            pdf.showPage()
            text_object = pdf.beginText(50, 800)
            text_object.setFont("Helvetica", 11)

        text_object.textLine(line[:100])

    pdf.drawText(text_object)
    pdf.save()

    pdf_file.seek(0)

    st.download_button(
        label="📕 Download PDF",
        data=pdf_file,
        file_name="LegalEase_Document.pdf",
        mime="application/pdf"
    )